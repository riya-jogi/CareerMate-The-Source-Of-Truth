import hashlib
import re
import uuid
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from app.core.config import settings
from app.core.errors import ConflictError, NotFoundError, ValidationError
from app.models.career import CandidateSkill, CareerClaim, ClaimStatus
from app.models.job import Job, JobRequirement
from app.models.profile import CareerProfile
from app.models.resume_version import Approval, ApprovalDecision, ChangeType, Resume, ResumeChange, ResumeVersion, ResumeVersionStatus, ValidationStatus
from app.models.user import User
from app.resume_generation_document import ResumeDocument, ResumeEducation, ResumeExperience, ResumeHeader, ResumeProject, ResumeCertification, format_date


class ResumeGenerationService:
    def _profile(self, db: Session, user: User) -> CareerProfile:
        profile = db.execute(
            select(CareerProfile)
            .options(
                selectinload(CareerProfile.experiences),
                selectinload(CareerProfile.projects),
                selectinload(CareerProfile.education),
                selectinload(CareerProfile.certifications),
                selectinload(CareerProfile.candidate_skills).selectinload(CandidateSkill.skill),
                selectinload(CareerProfile.career_claims),
            )
            .where(CareerProfile.user_id == user.id)
        ).scalar_one_or_none()
        if not profile:
            raise NotFoundError("Career profile was not found.")
        return profile

    def _job(self, db: Session, user: User, job_id: uuid.UUID) -> Job:
        job = db.execute(select(Job).where(Job.id == job_id, Job.user_id == user.id)).scalar_one_or_none()
        if not job:
            raise NotFoundError("Job was not found.")
        return job

    def build_document(self, profile: CareerProfile, user: User, job: Job | None = None) -> ResumeDocument:
        requirements = {self._normalize(item.skill_name) for item in (job.requirements if job else []) if item.skill_name}
        skills = [item.skill.name for item in profile.candidate_skills]
        relevant = [skill for skill in skills if self._normalize(skill) in requirements]
        ordered_skills = relevant + [skill for skill in skills if skill not in relevant]
        experiences = []
        for item in profile.experiences:
            bullets = [line.strip(" -*") for line in (item.description or "").splitlines() if line.strip()]
            if not bullets:
                bullets = [f"{item.job_title} at {item.company_name}"]
            experiences.append(ResumeExperience(source_id=item.id, company_name=item.company_name, job_title=item.job_title, location=item.location, start_date=item.start_date, end_date=item.end_date, bullets=bullets))
        return ResumeDocument(
            header=ResumeHeader(name=user.full_name, headline=profile.headline, email=user.email, phone=profile.phone, location=profile.location, links=[link for link in [profile.linkedin_url, profile.github_url, profile.portfolio_url] if link]),
            summary=profile.summary,
            skills=ordered_skills,
            experiences=experiences,
            projects=[ResumeProject(source_id=item.id, name=item.name, description=item.description, role=item.role) for item in profile.projects],
            education=[ResumeEducation(source_id=item.id, institution=item.institution, degree=item.degree, field_of_study=item.field_of_study, start_date=item.start_date, end_date=item.end_date) for item in profile.education],
            certifications=[ResumeCertification(source_id=item.id, name=item.name, issuing_organization=item.issuing_organization, issue_date=item.issue_date, expiry_date=item.expiry_date) for item in profile.certifications],
        )

    @staticmethod
    def _normalize(value: str | None) -> str:
        return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()

    def _get_resume(self, db: Session, profile_id: uuid.UUID, job_id: uuid.UUID | None = None) -> Resume | None:
        query = select(Resume).where(Resume.career_profile_id == profile_id).order_by(Resume.created_at)
        if job_id:
            query = query.options(selectinload(Resume.versions))
        return db.execute(query).scalars().first()

    def optimize(self, db: Session, user: User, job_id: uuid.UUID, template: str) -> ResumeVersion:
        profile = self._profile(db, user)
        job = self._job(db, user, job_id)
        resume = self._get_resume(db, profile.id)
        if not resume:
            resume = Resume(career_profile_id=profile.id, name="CareerMate Master Resume")
            db.add(resume)
            db.flush()
        version_number = (db.scalar(select(func.max(ResumeVersion.version_number)).where(ResumeVersion.resume_id == resume.id)) or 0) + 1
        document = self.build_document(profile, user, job)
        version = ResumeVersion(resume_id=resume.id, job_id=job.id, version_number=version_number, content_snapshot=document.model_dump(mode="json"), template=template, status=ResumeVersionStatus.PENDING_REVIEW, validation_status=ValidationStatus.PASS)
        db.add(version)
        db.flush()
        requirement_names = {self._normalize(item.skill_name): item for item in job.requirements if item.skill_name}
        for claim in profile.career_claims:
            subject = self._normalize(claim.skill_name or claim.subject)
            if subject in requirement_names and claim.status in {ClaimStatus.EVIDENCE_BACKED, ClaimStatus.CANDIDATE_CONFIRMED, ClaimStatus.SELF_DECLARED} and subject in {self._normalize(skill) for skill in document.skills}:
                db.add(ResumeChange(resume_version_id=version.id, change_type=ChangeType.REORDER, section="SKILLS", original_content="Candidate profile order", proposed_content=claim.subject, reason=f"Highlights {claim.subject} because it appears in the analyzed job requirements.", risk_level="low", validation_status=ValidationStatus.PASS, primary_claim_id=claim.id))
        db.flush()
        return self._load_version(db, version.id)

    def _load_version(self, db: Session, version_id: uuid.UUID) -> ResumeVersion:
        version = db.execute(
            select(ResumeVersion)
            .options(selectinload(ResumeVersion.changes).selectinload(ResumeChange.approval))
            .where(ResumeVersion.id == version_id)
        ).scalar_one_or_none()
        if not version:
            raise NotFoundError("Resume version was not found.")
        return version

    def decide_change(self, db: Session, user: User, change_id: uuid.UUID, decision: str, edited_text: str | None, comment: str | None) -> ResumeChange:
        change = db.execute(select(ResumeChange).join(ResumeVersion).join(Resume).join(CareerProfile).where(ResumeChange.id == change_id, CareerProfile.user_id == user.id)).scalar_one_or_none()
        if not change:
            raise NotFoundError("Resume change was not found.")
        if decision == "edited" and not edited_text:
            raise ValidationError("Edited decisions require edited_text.", {"code": "EDIT_TEXT_REQUIRED"})
        approval = change.approval or Approval(resume_change_id=change.id)
        approval.decision = ApprovalDecision(decision)
        approval.edited_text = edited_text
        approval.candidate_comment = comment
        db.add(approval)
        db.flush()
        # Reload with approval relationship populated
        return db.execute(
            select(ResumeChange).options(selectinload(ResumeChange.approval)).where(ResumeChange.id == change_id)
        ).scalar_one()

    def _approved_version(self, db: Session, user: User, version_id: uuid.UUID) -> ResumeVersion:
        version = db.execute(select(ResumeVersion).join(Resume).join(CareerProfile).options(selectinload(ResumeVersion.changes).selectinload(ResumeChange.approval)).where(ResumeVersion.id == version_id, CareerProfile.user_id == user.id)).scalar_one_or_none()
        if not version:
            raise NotFoundError("Resume version was not found.")
        pending = [change for change in version.changes if not change.approval]
        if pending:
            raise ConflictError("Every proposed change must be approved, rejected, or edited before generation.", {"pending_changes": len(pending)})
        blocked = [change for change in version.changes if change.validation_status == ValidationStatus.BLOCK]
        if blocked:
            raise ConflictError("The resume contains blocked changes and cannot be generated.", {"blocked_changes": len(blocked)})
        return version

    def generate(self, db: Session, user: User, version_id: uuid.UUID, output_format: str) -> tuple[ResumeVersion, Path]:
        version = self._approved_version(db, user, version_id)
        document = ResumeDocument.model_validate(version.content_snapshot)
        quality = self.validate_document(document)
        if quality["status"] != "ready":
            version.status = ResumeVersionStatus.NEEDS_REVIEW
            version.validation_status = ValidationStatus.NEEDS_CLARIFICATION
            db.flush()
            raise ValidationError("Resume quality validation did not pass.", quality)
        root = Path(settings.STORAGE_PATH).resolve() / "generated" / str(user.id) / str(version.id)
        root.mkdir(parents=True, exist_ok=True)
        path = root / f"resume.{output_format}"
        if output_format == "docx":
            self._render_docx(document, path)
            version.docx_key = str(path.relative_to(Path(settings.STORAGE_PATH).resolve()))
            version.docx_checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        else:
            self._render_pdf(document, path)
            version.pdf_key = str(path.relative_to(Path(settings.STORAGE_PATH).resolve()))
            version.pdf_checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        version.status = ResumeVersionStatus.READY
        version.validation_status = ValidationStatus.PASS
        db.flush()
        return version, path

    @staticmethod
    def validate_document(document: ResumeDocument) -> dict:
        issues: list[str] = []
        checks = {"content": "pass", "sections": "pass", "traceability": "pass", "format": "pass"}
        if not document.header.name or not document.header.email:
            issues.append("Header is missing candidate name or email.")
        if not document.experiences and not document.projects and not document.education:
            issues.append("Resume has no career content.")
        if any(len(bullet) > 500 for experience in document.experiences for bullet in experience.bullets):
            issues.append("One or more experience bullets are excessively long.")
        if issues:
            checks["content"] = "needs_review"
        return {"status": "ready" if not issues else "needs_review", "checks": checks, "issues": issues}

    @staticmethod
    def _render_docx(document: ResumeDocument, path: Path) -> None:
        output = Document()
        title = output.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(document.header.name)
        run.bold = True
        run.font.size = __import__("docx").shared.Pt(18)
        if document.header.headline:
            subtitle = output.add_paragraph(document.header.headline)
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        output.add_paragraph(" | ".join([document.header.email, document.header.phone or "", document.header.location or ""])).alignment = WD_ALIGN_PARAGRAPH.CENTER
        ResumeGenerationService._docx_section(output, "SUMMARY", [document.summary] if document.summary else [])
        ResumeGenerationService._docx_section(output, "SKILLS", [", ".join(document.skills)] if document.skills else [])
        ResumeGenerationService._docx_section(output, "EXPERIENCE", [f"{item.job_title} | {item.company_name} | {format_date(item.start_date)} - {format_date(item.end_date)}\n" + "\n".join(f"• {bullet}" for bullet in item.bullets) for item in document.experiences])
        ResumeGenerationService._docx_section(output, "PROJECTS", [f"{item.name}\n{item.description}" for item in document.projects])
        ResumeGenerationService._docx_section(output, "EDUCATION", [f"{item.degree}, {item.institution}" for item in document.education])
        ResumeGenerationService._docx_section(output, "CERTIFICATIONS", [f"{item.name}, {item.issuing_organization}" for item in document.certifications])
        output.save(path)

    @staticmethod
    def _docx_section(document, heading: str, paragraphs: list[str]) -> None:
        if not paragraphs:
            return
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)

    @staticmethod
    def _render_pdf(document: ResumeDocument, path: Path) -> None:
        styles = getSampleStyleSheet()
        story = [Paragraph(document.header.name, styles["Title"]), Paragraph(document.header.headline or "", styles["Normal"]), Paragraph(" | ".join([document.header.email, document.header.phone or "", document.header.location or ""]), styles["Normal"]), Spacer(1, 0.15 * inch)]
        sections = [("SUMMARY", [document.summary] if document.summary else []), ("SKILLS", [", ".join(document.skills)] if document.skills else []), ("EXPERIENCE", [f"{item.job_title} | {item.company_name} | {format_date(item.start_date)} - {format_date(item.end_date)}<br/>" + "<br/>".join(f"• {bullet}" for bullet in item.bullets) for item in document.experiences]), ("PROJECTS", [f"{item.name}<br/>{item.description}" for item in document.projects]), ("EDUCATION", [f"{item.degree}, {item.institution}" for item in document.education]), ("CERTIFICATIONS", [f"{item.name}, {item.issuing_organization}" for item in document.certifications])]
        for heading, paragraphs in sections:
            if paragraphs:
                story.append(Paragraph(heading, styles["Heading2"]))
                story.extend(Paragraph(value, styles["BodyText"]) for value in paragraphs)
                story.append(Spacer(1, 0.08 * inch))
        SimpleDocTemplate(str(path), pagesize=LETTER, rightMargin=0.65 * inch, leftMargin=0.65 * inch, topMargin=0.55 * inch, bottomMargin=0.55 * inch).build(story)


resume_generation_service = ResumeGenerationService()
